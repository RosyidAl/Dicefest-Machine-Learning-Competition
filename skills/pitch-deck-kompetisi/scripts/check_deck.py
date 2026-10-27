#!/usr/bin/env python3
"""Check a pitch-deck brief (.md) or deck (.pptx) against the rules of the pitch-deck-kompetisi skill.

Checks: em-dash anywhere (error), numbers in the slides that do not appear in the fact sheet(s) (error), on-screen word
count outside tables and table size per slide (warning), "kami"/"saya" on screen (warning), generic buzzwords
(warning), total spoken time against --minutes (warning).

Brief format: one heading per slide, "### Slide <n>. <title> (<x> menit <y> detik)"; speaker notes in a bullet that
starts with "- **Catatan pembicara:**" or "- **Narasi pembicara (...):**"; bullets that start with "- **Gambar:**" or
"- **Visual...:**" describe visuals and are not counted as on-screen text; "- **Istilah:**", "- **Bagian rubrik:**",
"- **Sumber:**" and "- **Pedoman jawaban:**" are metadata (not counted). Any other heading ends the current slide.

usage: python check_deck.py BRIEF.md|deck.pptx [--facts FILE ...] [--minutes 10] [--max-words 45]
"""
import argparse
import re
import subprocess
import sys
from pathlib import Path

EM_DASH = "—"
SLIDE_RE = re.compile(r"^#{2,4}\s*Slide\s+(\d+)\b[.:]?\s*(.*)$", re.I)
HEADING_RE = re.compile(r"^#{1,6}\s")
DUR_RE = re.compile(r"\((?:\s*(\d+)\s*menit)?(?:\s*(\d+)\s*detik)?\s*\)\s*$")
BULLET_RE = re.compile(r"^\s*(?:-|\d+\.)\s+")
NOTES_RE = re.compile(r"^\s*-\s*\*\*(Catatan|Narasi) pembicara[^*]*\*\*", re.I)
META_RE = re.compile(r"^\s*-\s*\*\*(Istilah|Bagian rubrik|Sumber|Pedoman jawaban)\b", re.I)
VISUAL_RE = re.compile(r"^\s*-\s*\*\*(Gambar|Visual)\b", re.I)
LABEL_RE = re.compile(r"\*\*[^*]{1,40}:\*\*")
NUM_RE = re.compile(r"(?<![\w.,/])[+-]?\d+(?:[.,]\d+)*")
PRONOUN_RE = re.compile(r"\b(kami|saya)\b", re.I)
# generic phrases that make a deck read as machine-written; flagged as warnings
BUZZ_RE = re.compile(r"\b(di era digital|revolusioner|holistik|sinergi\w*|inovatif|canggih|mutakhir|terdepan|"
                     r"luar biasa|game.?changer|solusi terbaik|tak tertandingi)\b", re.I)


def candidates(token):
    """all plain readings of a number token: Indonesian (1.234,5) and English (1,234.5 or 0.5) conventions; a
    separator is read as a thousands separator only when every group after it has exactly three digits"""
    t = token.lstrip("+-")
    out = {t}

    def thousands(sep):
        return all(len(g) == 3 for g in t.split(sep)[1:])

    if "," in t and "." in t:
        out.add(t.replace(".", "").replace(",", "."))
        out.add(t.replace(",", ""))
    elif "," in t:
        out.add(t.replace(",", "."))
        if thousands(","):
            out.add(t.replace(",", ""))
    elif "." in t and thousands("."):
        out.add(t.replace(".", ""))
    return out


def number_set(text):
    s = set()
    for tok in NUM_RE.findall(text):
        s |= candidates(tok)
    return s


def read_text(path):
    p = Path(path)
    if p.suffix.lower() == ".pdf":
        return subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True, check=True).stdout
    if p.suffix.lower() == ".pptx":
        return "\n".join(s["screen"] + "\n" + s["notes"] for s in pptx_slides(p))
    return p.read_text(encoding="utf-8")


def md_slides(text):
    """split a brief into slides with on-screen text, speaker notes and planned seconds"""
    slides, cur, mode = [], None, "screen"
    for line in text.splitlines():
        m = SLIDE_RE.match(line)
        if m:
            d = DUR_RE.search(m.group(2))
            secs = (int(d.group(1) or 0) * 60 + int(d.group(2) or 0)) if d and (d.group(1) or d.group(2)) else None
            cur = {"id": m.group(1), "title": m.group(2), "screen": "", "table": "", "notes": "", "seconds": secs,
                   "table_rows": 0}
            slides.append(cur)
            mode = "screen"
            continue
        if HEADING_RE.match(line):
            cur = None
            continue
        if cur is None:
            continue
        if NOTES_RE.match(line):
            mode = "notes"
            cur["notes"] += LABEL_RE.sub("", line) + "\n"
            continue
        if VISUAL_RE.match(line):
            mode = "visual"
            continue
        if META_RE.match(line):   # glossary / rubric tags / answer guides: shown small or not at all
            mode = "meta"
            continue
        if BULLET_RE.match(line) and not line.startswith((" ", "\t")):
            mode = "screen"
        if mode == "notes":
            cur["notes"] += line + "\n"
        elif mode == "screen":
            if line.strip().startswith("|"):   # tables are checked by size, not by word count
                if not set(line.replace("|", "").strip()) <= set("-: "):
                    cur["table_rows"] += 1
                    cur["table"] += line.replace("|", " ") + "\n"
            else:
                cur["screen"] += LABEL_RE.sub("", line) + "\n"
    return slides


def pptx_slides(path):
    try:
        from pptx import Presentation
    except ImportError:
        sys.exit("python-pptx is not installed: pip install python-pptx, or check the markdown brief instead")
    slides = []
    for i, s in enumerate(Presentation(str(path)).slides, 1):
        screen = []
        rows = 0
        for sh in s.shapes:
            if sh.has_text_frame:
                screen.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                rows += len(sh.table.rows)
                screen.extend(c.text for r in sh.table.rows for c in r.cells)
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        slides.append({"id": str(i), "title": screen[0][:60] if screen else "", "screen": "\n".join(screen),
                       "table": "", "notes": notes, "seconds": None, "table_rows": rows})
    return slides


def words(text):
    """words a viewer has to read: bullet markers, markdown syntax and lone symbols do not count"""
    text = re.sub(r"(?m)^\s*(?:-|\d+\.)\s+", " ", re.sub(r"[`*#]", " ", text))
    return [w for w in re.findall(r"[0-9A-Za-zÀ-ÿ²³⁻¹⁰%,.+/-]+", text) if re.search(r"[0-9A-Za-zÀ-ÿ]", w)]


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("deck")
    ap.add_argument("--facts", action="append", default=[], help="fact sheet, writeup (.md/.txt/.pdf); repeatable")
    ap.add_argument("--minutes", type=float, default=None, help="planned talk length")
    ap.add_argument("--max-words", type=int, default=45, help="on-screen words per slide before a warning")
    ap.add_argument("--max-table-rows", type=int, default=5)
    ap.add_argument("--min-number", type=float, default=13, help="integers below this are not checked (counts)")
    a = ap.parse_args()

    path = Path(a.deck)
    raw = read_text(path)
    slides = pptx_slides(path) if path.suffix.lower() == ".pptx" else md_slides(raw)
    errors, warnings = [], []

    for n, line in enumerate(raw.splitlines(), 1):
        if EM_DASH in line:
            errors.append(f"em-dash on line {n}: {line.strip()[:80]}")

    facts = set()
    for f in a.facts:
        facts |= number_set(read_text(f))

    for s in slides:
        tag = f"slide {s['id']}"
        w = len(words(s["screen"]))
        if w > a.max_words:
            warnings.append(f"{tag}: {w} on-screen words (> {a.max_words})")
        if s["table_rows"] > a.max_table_rows:
            warnings.append(f"{tag}: table with {s['table_rows']} rows (> {a.max_table_rows})")
        for m in sorted(set(PRONOUN_RE.findall(s["screen"].lower()))):
            warnings.append(f"{tag}: '{m}' in on-screen text")
        for m in sorted(set(x.lower() for x in BUZZ_RE.findall(s["screen"] + "\n" + s["notes"]))):
            warnings.append(f"{tag}: generic phrase '{m}'")
        if a.facts:
            missing = []
            for tok in NUM_RE.findall(s["screen"] + "\n" + s["table"] + "\n" + s["notes"]):
                c = candidates(tok)
                plain = tok.lstrip("+-")
                if re.fullmatch(r"\d+", plain) and int(plain) < a.min_number:
                    continue
                if not c & facts:
                    missing.append(tok)
            if missing:
                errors.append(f"{tag}: numbers not in the fact sheet: {', '.join(sorted(set(missing)))}")

    timed = [s["seconds"] for s in slides if s["seconds"] is not None]
    if timed:
        total = sum(timed)
        line = f"planned time {total // 60} min {total % 60} s over {len(timed)} slides"
        if a.minutes and abs(total - a.minutes * 60) > 0.1 * a.minutes * 60:
            warnings.append(line + f" (target {a.minutes:g} min)")
        else:
            print(line)

    print(f"{len(slides)} slides checked, {len(errors)} errors, {len(warnings)} warnings")
    for e in errors:
        print("ERROR  ", e)
    for w in warnings:
        print("WARNING", w)
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
