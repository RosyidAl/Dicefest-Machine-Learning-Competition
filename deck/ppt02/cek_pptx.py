"""Checks BISMILLAH_Final_v3.pptx against its source, 01_struktur_deck_minto.md (through build/isi_v3.json).

Main slides: the title equals the markdown title; every word and number of "Isi Teks" (table cells included) is on
the slide; the speaker notes hold the timing line and the full narration; every number on the slide appears in the
markdown slide (title, visual, on-screen text, narration) or in the chart data (img/data_grafik.json).
Backup slides: title, the "CADANGAN ... Cn" kicker, every table cell of the evidence, the questions and the answer
guide in the notes, and every number on the slide present in the markdown section of that backup.
Whole deck: 13 + 11 slides in order, the six sections, no em-dash, no "kami"/"saya" on screen, no font below 10 pt,
every shape inside the slide, and the speaker-name placeholders still to be filled.

usage: <python with python-pptx> deck/ppt02/cek_pptx.py [deck.pptx]
"""
import json
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
ISI = json.loads((HERE / "build" / "isi_v3.json").read_text(encoding="utf-8"))
DATA = json.loads((HERE / "img" / "data_grafik.json").read_text(encoding="utf-8"))
DECK = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "BISMILLAH_Final_v3.pptx"
SECTIONS = [("I. Pembukaan", 2), ("II. Rekayasa fitur", 3), ("III. Arsitektur model", 2), ("IV. Kebijakan publik", 4),
            ("V. Keterbatasan dan penutup", 2), ("Cadangan tanya jawab", 11)]
EM_DASH = "\u2014"
WORD = re.compile(r"[0-9a-zà-ÿ²_]+(?:[.,][0-9]+)*")
NUM = re.compile(r"(?<![\w.,])\d+(?:[.,]\d+)*")


def norm(t):
    """line breaks (soft or hard) and non-breaking spaces read as one space"""
    return re.sub(r"\s+", " ", t.replace("\u00a0", " ").replace("\v", " ")).strip()


def plain(t):
    return re.sub(r"`([^`]+)`", r"\1", re.sub(r"\*([^*]+)\*", r"\1", re.sub(r"\*\*([^*]+)\*\*", r"\1", t)))


def tokens(t):
    return set(WORD.findall(plain(t).lower().replace("−", "-")))


def numbers(t):
    """numbers in a text; bare integers up to 9 (step numbers, ranks, counts) are left out"""
    return {n.rstrip(".,") for n in NUM.findall(plain(t)) if not re.fullmatch(r"\d", n)}


def block_text(blocks):
    out = []
    for b in blocks:
        out.extend(c for row in b["baris"] for c in row) if b["jenis"] == "tabel" else out.append(b["teks"])
    return "\n".join(out)


def shapes_text(slide):
    """all on-screen text of a slide: text frames (placeholders included) and table cells"""
    parts = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:  # group
                walk(sh.shapes)
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts.extend(c.text for r in sh.table.rows for c in r.cells)
    walk(slide.shapes)
    return "\n".join(parts)


def chart_numbers():
    """numbers drawn from the chart data, formatted as on the slides"""
    out = set()

    def fmt(x, nd):
        return f"{x:.{nd}f}".replace(".", ",")
    for rows in DATA["s3_shap_top"].values():
        out |= {fmt(r["nilai"], 4) for r in rows}
    for r in DATA["s9_retrofit"]:
        out |= {fmt(r["gwh"], 0), fmt(r["pct"], 2)}
    k = DATA["s10_kenaikan"]
    out |= {fmt(r["pct"], 2) for r in k["lima"]} | {str(r["peringkat_cop_iklim"]) for r in k["lima"]}
    out |= {fmt(k["median"], 2), fmt(k["tambahan_absolut_terbesar"]["TWh"], 2)}
    for r in DATA["s11_adopsi"]:
        out |= {fmt(r["adopsi"], 2), fmt(r["kawasan"], 2), *numbers(r["pita"])}
    for t in DATA["c5_kelompok"]["target"]:
        out |= {str(p) for p in t["porsi"]}
    return out


def main():
    prs = Presentation(str(DECK))
    sw, sh = prs.slide_width, prs.slide_height
    slides = list(prs.slides)
    errors, notes_todo, small = [], [], 0
    if len(slides) != len(ISI["utama"]) + len(ISI["cadangan"]):
        errors.append(f"{len(slides)} slides, expected {len(ISI['utama']) + len(ISI['cadangan'])}")
    charts = chart_numbers()

    for idx, s in enumerate(slides, 1):
        text = shapes_text(s)
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        # pptxgenjs gives the title placeholder idx 101 (python-pptx's shapes.title only looks at idx 0)
        title = next((ph.text_frame.text for ph in s.placeholders if ph.placeholder_format.type == 1), "")
        tag = f"slide {idx}"
        if EM_DASH in text or EM_DASH in notes:
            errors.append(f"{tag}: em-dash")
        for m in re.findall(r"\b(kami|saya)\b", text.lower()):
            errors.append(f"{tag}: '{m}' on screen")
        if "[nama Orang" in notes:
            notes_todo.append(idx)
        # fonts and bounds
        for shp in s.shapes:
            if shp.left is not None and (shp.left < -Emu(9144) or shp.top < -Emu(9144) or
                                         shp.left + shp.width > sw + Emu(9144) or shp.top + shp.height > sh + Emu(9144)):
                errors.append(f"{tag}: '{shp.name}' outside the slide")
            frames = [shp.text_frame] if shp.has_text_frame else []
            if getattr(shp, "has_table", False) and shp.has_table:
                frames += [c.text_frame for r in shp.table.rows for c in r.cells]
            for tf in frames:
                for p in tf.paragraphs:
                    for r in p.runs:
                        if r.font.size is not None and r.font.size.pt < 10:
                            errors.append(f"{tag}: {r.font.size.pt} pt text '{r.text[:30]}'")
                        elif r.font.size is not None and r.font.size.pt < 12 and r.text.strip():
                            small += 1

        if idx <= len(ISI["utama"]):
            d = ISI["utama"][idx - 1]
            if norm(title) != norm(plain(d["judul"])):
                errors.append(f"{tag}: title '{title[:60]}' differs from the markdown")
            missing = tokens(block_text(d["isi"])) - tokens(text)
            if missing:
                errors.append(f"{tag}: on-screen words missing: {sorted(missing)}")
            head = f"[{d['mulai']}–{d['selesai']} · {d['detik']} detik · {d['pembicara']}]"
            if not notes.startswith(head) or plain(d["narasi"]) not in notes:
                errors.append(f"{tag}: notes lack the timing line or the full narration")
            allowed = numbers(" ".join([d["judul"], d["visual"], block_text(d["isi"]), d["narasi"]])) | charts | {"2026"}
            extra = numbers(text) - allowed - {str(idx)}
            if extra:
                errors.append(f"{tag}: numbers not in the markdown slide: {sorted(extra)}")
        else:
            b = ISI["cadangan"][idx - len(ISI["utama"]) - 1]
            if norm(title) != norm(plain(b["judul"])):
                errors.append(f"{tag}: title differs from {b['id']} in the markdown")
            if f"CADANGAN TANYA JAWAB · {b['id']}" not in text:
                errors.append(f"{tag}: kicker without {b['id']}")
            cells = [c for blk in b["bukti"] if blk["jenis"] == "tabel" for row in blk["baris"] for c in row]
            missing = tokens("\n".join(cells)) - tokens(text)
            if missing:
                errors.append(f"{tag} ({b['id']}): table words missing: {sorted(missing)}")
            if plain(b["pedoman"]) not in notes or any(q not in notes for q in b["pertanyaan"]):
                errors.append(f"{tag} ({b['id']}): notes lack the questions or the answer guide")
            section = " ".join([b["judul"], block_text(b["bukti"]), b["pedoman"], " ".join(b["pertanyaan"])])
            allowed = numbers(section) | charts | {"2026", str(idx)}
            extra = numbers(text) - allowed
            if extra:
                errors.append(f"{tag} ({b['id']}): numbers not in the markdown section: {sorted(extra)}")

    # sections (PowerPoint 2010 extension list in presentation.xml)
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
          "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main"}
    root = etree.fromstring(prs.part.blob)
    found = [(sec.get("name"), len(sec.findall(".//p14:sldId", ns))) for sec in root.findall(".//p14:section", ns)]
    if found != SECTIONS:
        errors.append(f"sections {found} differ from {SECTIONS}")

    print(f"{DECK.name}: {len(slides)} slides, {len(errors)} errors; {small} text runs at 10 to 11 pt (labels, captions)")
    print(f"speaker-name placeholders still in the notes of slides {notes_todo}")
    for e in errors:
        print("ERROR ", e)
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
