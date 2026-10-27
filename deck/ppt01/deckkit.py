"""Tiny slide kit: one list of elements per slide, written to PPTX (python-pptx) and rendered to a PNG preview (PIL).

The preview uses Carlito, which has the same glyph metrics as Calibri (the font written into the PPTX), so text
wrapping and overflow seen in the preview match PowerPoint closely. Elements are dicts:
  text  {x, y, w, h, p: [paragraph, ...], size, color, bold, italic, align, valign, lsp, after, role}
        a paragraph is a string or {text, size, color, bold, align, after, bullet}; inline markup:
        **bold**, *italic*, {a|amber} {A|amber bold} {t|teal} {T|..} {n|navy} {N|..} {m|muted} {w|white} {W|..}
  rect  {x, y, w, h, fill, line, lw, r}      oval {x, y, w, h, fill, line}
  line  {x1, y1, x2, y2, color, lw, arrow}   img  {path, x, y, w, h, alt}
"""
import os
import re

from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
FONTS = {(False, False): "Carlito-Regular.ttf", (True, False): "Carlito-Bold.ttf",
         (False, True): "Carlito-Italic.ttf", (True, True): "Carlito-BoldItalic.ttf"}
SW, SH = 13.333, 7.5
LINE_H = 1.2207  # Calibri / Carlito: (ascender + descender + line gap) / em
C = dict(NAVY="16324F", TEAL="1F7A8C", AMBER="D97B29", INK="1F2933", MUTED="5F6B7A", LINE="D5DBE3", SOFT="EEF2F6",
         WHITE="FFFFFF", TEAL_T="E3F1F4", AMBER_T="FBEEDF", NAVY_T="E3EAF2", VIOLET="5B4A86", VIOLET_T="ECE8F4",
         NAVY2="22476A", LIGHT="C3D0DE", TEAL_D="17606E", AMBER_D="9A5212")
MARK = {"^": ("^", False), "a": ("AMBER", False), "A": ("AMBER", True), "t": ("TEAL", False), "T": ("TEAL", True), "n": ("NAVY", False),
        "N": ("NAVY", True), "m": ("MUTED", False), "w": ("WHITE", False), "W": ("WHITE", True),
        "v": ("VIOLET", False), "d": ("AMBER_D", True)}
TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|\{[a-zA-Z^]\|.+?\})")
_font_cache = {}


def col(c):
    return C.get(c, c)


def font(bold, italic, px):
    key = (bold, italic, round(px * 4))
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(os.path.join(HERE, "fonts", FONTS[(bold, italic)]), max(1, round(px)))
    return _font_cache[key]


def runs(text, bold=False, italic=False, color="INK"):
    """inline markup -> list of (text, bold, italic, color)"""
    out = []
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, italic, color))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], bold, True, color))
        elif part.startswith("{") and part.endswith("}") and part[2] == "|":
            c, b = MARK[part[1]]
            out.append((part[3:-1], b or bold, italic, "^" + color if c == "^" else c))  # "^" prefix = superscript
        else:
            out.append((part, bold, italic, color))
    return out


def para_list(e):
    ps = []
    for p in e["p"]:
        d = {"text": p} if isinstance(p, str) else dict(p)
        for k in ("size", "color", "bold", "italic", "align", "after", "lsp"):
            d.setdefault(k, e.get(k, {"size": 18, "color": "INK", "bold": False, "italic": False, "align": "l",
                                      "after": 0, "lsp": 1.0}[k]))
        d.setdefault("bullet", False)
        ps.append(d)
    return ps


def text_width_in(text, size, bold=False, italic=False):
    """width in inches of a single line of plain text"""
    f = font(bold, italic, size * 4)
    return f.getlength(text) / (size * 4) * size / 72


# ---------------------------------------------------------------------------------------------------------------
# layout of one paragraph into lines (shared by the preview and by the overflow check)
def run_px(px, c):
    return px * 0.6 if c.startswith("^") else px


def layout_para(d, width_in, px_per_in):
    px = d["size"] * px_per_in / 72
    indent = 0.25 * px_per_in if d["bullet"] else 0
    maxw = width_in * px_per_in - indent
    words, cur = [], []
    for t, b, i, c in runs(d["text"], d["bold"], d["italic"], d["color"]):
        for piece in re.findall(r"\s+|[^\s]+", t):
            if piece.isspace():
                if cur:
                    words.append(cur)
                    cur = []
                words.append([(" ", b, i, c)])
            else:
                cur.append((piece, b, i, c))
    if cur:
        words.append(cur)
    lines, line, lw, too_wide = [], [], 0.0, False
    for w in words:
        ww = sum(font(b, i, run_px(px, c)).getlength(s) for s, b, i, c in w)
        is_space = w[0][0] == " "
        if is_space:
            if line:
                line.append(w)
                lw += ww
            continue
        if line and lw + ww > maxw + 0.5:
            while line and line[-1][0][0] == " ":
                lw -= sum(font(b, i, run_px(px, c)).getlength(s) for s, b, i, c in line.pop())
            lines.append((line, lw))
            line, lw = [], 0.0
        if ww > maxw + 0.5:
            too_wide = True
        line.append(w)
        lw += ww
    while line and line[-1][0][0] == " ":
        lw -= sum(font(b, i, run_px(px, c)).getlength(s) for s, b, i, c in line.pop())
    lines.append((line, lw))
    line_h = px * LINE_H * d["lsp"]
    return lines, line_h, indent, too_wide


def text_height_in(e):
    total = 0.0
    ps = para_list(e)
    for k, d in enumerate(ps):
        lines, line_h, _, _ = layout_para(d, e["w"], 288)
        total += len(lines) * line_h / 288 + (d["after"] / 72 if k < len(ps) - 1 else 0)
    return total


def check_fit(slides):
    """text boxes whose wrapped text is taller than the box, or with a word wider than the box"""
    issues = []
    for n, s in enumerate(slides, 1):
        for e in s["el"]:
            if e["t"] != "text":
                continue
            need = text_height_in(e)
            wide = any(layout_para(d, e["w"], 288)[3] for d in para_list(e))
            if need > e["h"] + 0.02 or wide:
                first = para_list(e)[0]["text"][:50]
                issues.append(f"{s['id']}: '{first}' needs {need:.2f} in, box {e['h']:.2f} in"
                              + (" (word wider than box)" if wide else ""))
    return issues


# ---------------------------------------------------------------------------------------------------------------
# preview renderer
def hexrgb(c):
    c = col(c)
    return tuple(int(c[i:i + 2], 16) for i in (0, 2, 4))


def render(slide, path, dpi=144):
    W, H = round(SW * dpi), round(SH * dpi)
    im = Image.new("RGB", (W, H), hexrgb(slide.get("bg") or "WHITE"))
    dr = ImageDraw.Draw(im)
    P = lambda v: round(v * dpi)  # noqa: E731
    for e in slide["el"]:
        t = e["t"]
        if t in ("rect", "oval"):
            box = [P(e["x"]), P(e["y"]), P(e["x"] + e["w"]), P(e["y"] + e["h"])]
            fill = hexrgb(e["fill"]) if e.get("fill") else None
            outline = hexrgb(e["line"]) if e.get("line") else None
            width = max(1, round(e.get("lw", 1) * dpi / 72)) if outline else 0
            if t == "oval":
                dr.ellipse(box, fill=fill, outline=outline, width=width)
            elif e.get("r"):
                dr.rounded_rectangle(box, radius=P(e["r"]), fill=fill, outline=outline, width=width)
            else:
                dr.rectangle(box, fill=fill, outline=outline, width=width)
        elif t == "line":
            w = max(1, round(e.get("lw", 1.5) * dpi / 72))
            x1, y1, x2, y2 = P(e["x1"]), P(e["y1"]), P(e["x2"]), P(e["y2"])
            dr.line([x1, y1, x2, y2], fill=hexrgb(e.get("color", "MUTED")), width=w)
            if e.get("arrow"):
                import math
                ang = math.atan2(y2 - y1, x2 - x1)
                L, Wd = 0.13 * dpi, 0.07 * dpi
                p1 = (x2 - L * math.cos(ang) + Wd * math.sin(ang), y2 - L * math.sin(ang) - Wd * math.cos(ang))
                p2 = (x2 - L * math.cos(ang) - Wd * math.sin(ang), y2 - L * math.sin(ang) + Wd * math.cos(ang))
                dr.polygon([(x2, y2), p1, p2], fill=hexrgb(e.get("color", "MUTED")))
        elif t == "img":
            pic = Image.open(e["path"]).convert("RGB").resize((P(e["w"]), P(e["h"])), Image.LANCZOS)
            im.paste(pic, (P(e["x"]), P(e["y"])))
        elif t == "text":
            ps = para_list(e)
            laid = [layout_para(d, e["w"], dpi) for d in ps]
            total = sum(len(l[0]) * l[1] for l in laid) + sum(d["after"] * dpi / 72 for d in ps[:-1])
            y = P(e["y"])
            va = e.get("valign", "t")
            if va == "m":
                y += (P(e["h"]) - total) / 2
            elif va == "b":
                y += P(e["h"]) - total
            for d, (lines, line_h, indent, _) in zip(ps, laid):
                px = d["size"] * dpi / 72
                for k, (line, lw) in enumerate(lines):
                    x0 = P(e["x"]) + indent
                    if d["align"] == "c":
                        x0 = P(e["x"]) + (P(e["w"]) - lw) / 2
                    elif d["align"] == "r":
                        x0 = P(e["x"] + e["w"]) - lw
                    base = y + line_h * 0.5 + px * 0.31  # visually centred baseline in the line box
                    if d["bullet"] and k == 0:
                        dr.text((P(e["x"]), base), "•", font=font(False, False, px), fill=hexrgb(d["color"]),
                                anchor="ls")
                    x = x0
                    for w in line:
                        for s, b, i, c in w:
                            sup = c.startswith("^")
                            f = font(b, i, run_px(px, c))
                            dr.text((x, base - (px * 0.33 if sup else 0)), s, font=f, fill=hexrgb(c.lstrip("^")),
                                    anchor="ls")
                            x += f.getlength(s)
                    y += line_h
                y += d["after"] * dpi / 72
    im.save(path)
    return im


# ---------------------------------------------------------------------------------------------------------------
# PPTX writer
def write_pptx(slides, path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt
    from lxml import etree

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    layout = prs.slide_layouts[5]  # title only: the action title goes into the real title placeholder
    AL = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
    VA = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

    def drop_style(shape):
        st = shape._element.find(qn("p:style"))
        if st is not None:
            shape._element.remove(st)

    def fill_text(tf, e):
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = VA[e.get("valign", "t")]
        for k, d in enumerate(para_list(e)):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = AL[d["align"]]
            p.line_spacing = d["lsp"]
            p.space_after = Pt(d["after"])
            p.space_before = Pt(0)
            if d["bullet"]:
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(Inches(0.25)))
                pPr.set("indent", str(-Inches(0.25)))
                bf = etree.SubElement(pPr, qn("a:buFont"))
                bf.set("typeface", "Arial")
                bc = etree.SubElement(pPr, qn("a:buChar"))
                bc.set("char", "•")
            for s, b, i, c in runs(d["text"], d["bold"], d["italic"], d["color"]):
                r = p.add_run()
                r.text = s
                f = r.font
                f.name = "Calibri"
                f.size = Pt(d["size"])
                f.bold = b
                f.italic = i
                f.color.rgb = RGBColor.from_string(col(c.lstrip("^")))
                if c.startswith("^"):
                    r._r.get_or_add_rPr().set("baseline", "30000")
                if e.get("spacing"):
                    r._r.get_or_add_rPr().set("spc", str(int(e["spacing"] * 100)))

    for s in slides:
        sl = prs.slides.add_slide(layout)
        if s.get("bg"):
            sl.background.fill.solid()
            sl.background.fill.fore_color.rgb = RGBColor.from_string(col(s["bg"]))
        for e in s["el"]:
            t = e["t"]
            if t in ("rect", "oval"):
                kind = MSO_SHAPE.OVAL if t == "oval" else (MSO_SHAPE.ROUNDED_RECTANGLE if e.get("r") else
                                                          MSO_SHAPE.RECTANGLE)
                sh = sl.shapes.add_shape(kind, Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
                drop_style(sh)
                if e.get("r"):
                    sh.adjustments[0] = min(0.5, e["r"] / min(e["w"], e["h"]))
                if e.get("fill"):
                    sh.fill.solid()
                    sh.fill.fore_color.rgb = RGBColor.from_string(col(e["fill"]))
                else:
                    sh.fill.background()
                if e.get("line"):
                    sh.line.color.rgb = RGBColor.from_string(col(e["line"]))
                    sh.line.width = Pt(e.get("lw", 1))
                else:
                    sh.line.fill.background()
                if e.get("name"):
                    sh.name = e["name"]
            elif t == "line":
                cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(e["x1"]), Inches(e["y1"]),
                                             Inches(e["x2"]), Inches(e["y2"]))
                drop_style(cn)
                cn.line.color.rgb = RGBColor.from_string(col(e.get("color", "MUTED")))
                cn.line.width = Pt(e.get("lw", 1.5))
                if e.get("arrow"):
                    ln = cn.line._get_or_add_ln()
                    tail = etree.SubElement(ln, qn("a:tailEnd"))
                    tail.set("type", "triangle")
                    tail.set("w", "med")
                    tail.set("len", "med")
            elif t == "img":
                pic = sl.shapes.add_picture(e["path"], Inches(e["x"]), Inches(e["y"]), Inches(e["w"]),
                                            Inches(e["h"]))
                pic._element.nvPicPr.cNvPr.set("descr", e.get("alt", os.path.basename(e["path"])))
            elif t == "text":
                if e.get("role") == "title" and sl.shapes.title is not None and not sl.shapes.title.has_text_frame \
                        or (e.get("role") == "title" and sl.shapes.title is not None
                            and not sl.shapes.title.text_frame.text):
                    tb = sl.shapes.title
                    tb.left, tb.top, tb.width, tb.height = (Inches(e["x"]), Inches(e["y"]), Inches(e["w"]),
                                                            Inches(e["h"]))
                else:
                    tb = sl.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
                    if e.get("role"):
                        tb.name = e["role"]
                fill_text(tb.text_frame, e)
        if sl.shapes.title is not None and not sl.shapes.title.text_frame.text:  # no title element: drop placeholder
            sl.shapes.title._element.getparent().remove(sl.shapes.title._element)
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]
    prs.core_properties.title = slides[0].get("doc_title", "Deck")
    prs.core_properties.author = "Tim BISMILLAH"
    prs.save(path)
